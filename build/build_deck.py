# -*- coding: utf-8 -*-
import os
import re
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import copy

# ---------- app icon assets: pad to square (transparent) so a circle crop never clips them ----------
IMG_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "icons")
ICON_SRC = {
    "wechat": os.path.join(IMG_DIR, "2.png"),
    "weibo": os.path.join(IMG_DIR, "3.png"),
    "bilibili": os.path.join(IMG_DIR, "4.png"),
    "douyin": os.path.join(IMG_DIR, "5.png"),
    "xiaohongshu": os.path.join(IMG_DIR, "6.png"),
    "facebook": os.path.join(IMG_DIR, "13.png"),
    "instagram": os.path.join(IMG_DIR, "14.png"),
    "x_twitter": os.path.join(IMG_DIR, "15.png"),
    "kakaotalk": os.path.join(IMG_DIR, "16.png"),
    "youtube": os.path.join(IMG_DIR, "17.png"),
}
ICON_SQ_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "icons_sq")
os.makedirs(ICON_SQ_DIR, exist_ok=True)


def square_pad(src_path, key):
    im = Image.open(src_path).convert("RGBA")
    w, h = im.size
    s = max(w, h)
    canvas = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    canvas.paste(im, ((s - w) // 2, (s - h) // 2), im)
    out_path = os.path.join(ICON_SQ_DIR, key + ".png")
    canvas.save(out_path)
    return out_path


ICON = {key: square_pad(path, key) for key, path in ICON_SRC.items()}

# ---------- palette ----------
RED       = RGBColor(0xC8, 0x10, 0x2E)   # primary china red
DARK_RED  = RGBColor(0x3D, 0x08, 0x11)   # deep background red
GOLD      = RGBColor(0xC9, 0xA2, 0x27)   # accent gold
GLYPH_INK = RGBColor(0xF7, 0xF7, 0xF5)   # hanzi glyph color inside red badges
DARK      = RGBColor(0x22, 0x1F, 0x1F)   # near-black warm charcoal (body text)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_CARD = RGBColor(0xF3, 0xF3, 0xF3)   # light neutral card
MUTED     = RGBColor(0x8A, 0x83, 0x83)   # muted gray for de-emphasized text
LIGHT_TXT = RGBColor(0xEA, 0xE2, 0xDD)   # off-white text on dark bg

TITLE_FONT = "A2Z 4 Regular"
BODY_FONT  = "A2Z 4 Regular"
GLYPH_FONT = "Noto Sans SC Medium"
HEADLINE_FONT = "A2Z 7 Bold"   # 슬라이드별 큰 제목 전용 폰트

SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width  = Emu(int(SW * 914400))
prs.slide_height = Emu(int(SH * 914400))
BLANK = prs.slide_layouts[6]


def no_shadow(shp):
    shp.shadow.inherit = False


def add_rect(slide, x, y, w, h, fill, radius=None, line_color=None, line_w=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_color is not None:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(line_w or 0.75)
    else:
        shp.line.fill.background()
    no_shadow(shp)
    return shp


def add_circle(slide, cx, cy, d, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    no_shadow(shp)
    return shp


HANZI_RE = re.compile(r'([一-鿿]+)')


def is_hanzi_seg(seg):
    return bool(HANZI_RE.fullmatch(seg))


def add_run_mixed(p, text, size, color, bold=False, italic=False, font=BODY_FONT, hanzi_font=GLYPH_FONT):
    for seg in [s for s in HANZI_RE.split(text) if s]:
        r = p.add_run()
        r.text = seg
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = hanzi_font if is_hanzi_seg(seg) else font
        r.font.color.rgb = color


def add_text(slide, x, y, w, h, text, size, color, bold=False, italic=False,
             align=PP_ALIGN.LEFT, font=BODY_FONT, anchor=None, wrap=True, hanzi_font=GLYPH_FONT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if anchor is not None:
        tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        add_run_mixed(p, line, size, color, bold=bold, italic=italic, font=font, hanzi_font=hanzi_font)
    return tb


def add_badge(slide, cx, cy, d, glyph, glyph_size, fill=RED, glyph_color=GLYPH_INK):
    add_circle(slide, cx, cy, d, fill)
    add_text(slide, cx - d / 2, cy - d / 2, d, d, glyph, glyph_size, glyph_color,
              bold=True, align=PP_ALIGN.CENTER, font=GLYPH_FONT,
              anchor=MSO_ANCHOR.MIDDLE, wrap=False)


def crop_to_ellipse(picture):
    geom = picture._element.spPr.find(qn('a:prstGeom'))
    if geom is not None:
        geom.set('prst', 'ellipse')


def add_icon_circle(slide, image_key, cx, cy, d, img_scale=0.82):
    plate = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    plate.fill.solid()
    plate.fill.fore_color.rgb = WHITE
    plate.line.color.rgb = RGBColor(0xE6, 0xE1, 0xDF)
    plate.line.width = Pt(1)
    no_shadow(plate)
    img_d = d * img_scale
    pic = slide.shapes.add_picture(ICON[image_key], Inches(cx - img_d / 2), Inches(cy - img_d / 2),
                                     Inches(img_d), Inches(img_d))
    crop_to_ellipse(pic)
    return plate


def add_bullets(slide, x, y, w, h, items, size, color, bullet_color=RED,
                 space_after=8, font=BODY_FONT, line_spacing=1.08):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        r1 = p.add_run()
        r1.text = "•  "
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.name = font
        r1.font.color.rgb = bullet_color
        add_run_mixed(p, item, size, color, font=font)
    return tb


def add_pill(slide, x, y, w, h, text, fill, text_color, size=11):
    add_rect(slide, x, y, w, h, fill, radius=0.5)
    add_text(slide, x, y, w, h, text, size, text_color, bold=True,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)


def add_gradient_rect(slide, x, y, w, h, color_from, color_to, angle=45):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.line.fill.background()
    no_shadow(shp)
    shp.fill.gradient()
    stops = shp.fill.gradient_stops
    stops[0].position = 0.0
    stops[0].color.rgb = color_from
    stops[1].position = 1.0
    stops[1].color.rgb = color_to
    shp.fill.gradient_angle = angle
    return shp


def add_bubble(slide, x, y, w, h, glyph, fill=GOLD, glyph_size=18):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = -0.32
        shp.adjustments[1] = 0.9
        shp.adjustments[2] = 0.28
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    no_shadow(shp)
    tf = shp.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = glyph
    r.font.size = Pt(glyph_size)
    r.font.color.rgb = DARK
    return shp


def add_vertical_label(slide, cx, cy, w, h, text, color, size=11, rotation=-90):
    tb = slide.shapes.add_textbox(Inches(cx - w / 2), Inches(cy - h / 2), Inches(w), Inches(h))
    tb.rotation = rotation
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = BODY_FONT
    r.font.color.rgb = color
    return tb


def add_heading(slide, kicker, title, title_color=RED, title_size=30):
    add_text(slide, 0.7, 0.42, 8.0, 0.35, kicker, 13, GOLD, bold=True, font=BODY_FONT)
    add_text(slide, 0.7, 0.75, 10.5, 0.7, title, title_size, title_color, bold=False, font=HEADLINE_FONT)


def add_highlight_subtitle(slide, prefix, highlight, suffix, y, size=15.5,
                            base_color=DARK, hl_color=RED, x=0.7, w=None, h=0.55):
    if w is None:
        w = SW - 2 * x
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run_mixed(p, prefix, size, base_color, font=BODY_FONT)
    add_run_mixed(p, highlight, size, hl_color, font=BODY_FONT)
    add_run_mixed(p, suffix, size, base_color, font=BODY_FONT)
    return tb


def add_page_number(slide, n):
    add_text(slide, SW - 1.2, SH - 0.5, 0.8, 0.3, f"{n} / 5", 10, MUTED,
              align=PP_ALIGN.RIGHT, wrap=False)


# =====================================================================
# Slide 1 — Title
# =====================================================================
s = prs.slides.add_slide(BLANK)
ICON_SOFT = RGBColor(0xF0, 0xC9, 0xC0)
DOT_FAINT = RGBColor(0xF3, 0xDD, 0xD8)

add_rect(s, 0, 0, SW, SH, WHITE)


def add_dot_grid(slide, x0, y0, cols, rows, spacing, radius, color):
    for r in range(rows):
        for c in range(cols):
            if (r + c) % 4 == 3:
                continue
            add_circle(slide, x0 + c * spacing, y0 + r * spacing, radius, color)


# faint corner texture, drawn first so the phone card covers most of it
add_dot_grid(s, 0.35, 0.35, 7, 5, 0.42, 0.045, DOT_FAINT)
add_dot_grid(s, SW - 0.35 - 6 * 0.42, SH - 0.35 - 4 * 0.42, 7, 5, 0.42, 0.045, DOT_FAINT)

# vertical decorative labels in the side margins, outside the phone card

# ---- phone mockup card, centered ----
px, py, pw, ph = 0.75, 0.6, SW - 1.5, SH - 1.2

card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(px), Inches(py), Inches(pw), Inches(ph))
card.adjustments[0] = 0.055
card.line.color.rgb = DARK
card.line.width = Pt(3.25)
no_shadow(card)
card.fill.solid()
card.fill.fore_color.rgb = RED

# phone notch, straddling the top bezel edge
notch_w, notch_h = 2.0, 0.22
notch = add_rect(s, SW / 2 - notch_w / 2, py - notch_h / 2, notch_w, notch_h, DARK, radius=0.5)

# side buttons, protruding from the left bezel edge
add_rect(s, px - 0.09, py + 0.85, 0.14, 0.55, DARK, radius=0.5)
add_rect(s, px - 0.09, py + 1.55, 0.14, 0.55, DARK, radius=0.5)

# top status bar (inside the phone screen)
add_text(s, px + 0.35, py + 0.28, 0.6, 0.4, "☰", 17, ICON_SOFT, bold=True, font=BODY_FONT, wrap=False)
add_text(s, px + pw - 2.5, py + 0.28, 2.15, 0.4, "\U0001F4F6   \U0001F50B   ▦", 13, ICON_SOFT,
          align=PP_ALIGN.RIGHT, font=BODY_FONT, wrap=False)

# kicker (centered) + chat bubble accent to its upper right
add_text(s, px, py + 1.0, pw, 0.5, "CHINA SOCIAL MEDIA", 13.5, GLYPH_INK, bold=True,
          align=PP_ALIGN.CENTER, font=BODY_FONT, wrap=False)
add_bubble(s, SW / 2 + 2.55, py + 0.88, 0.8, 0.56, "\U0001F44D", fill=GLYPH_INK, glyph_size=18)
add_bubble(s, SW / 2 + 3.25, py + 1.1, 0.55, 0.4, "❤️", fill=GLYPH_INK, glyph_size=13)

# big title, one line, centered
add_text(s, px + 0.37, py + 1.7, pw - 0.74, 1.1, "중국 SNS 완전정복", 40, WHITE, bold=False,
          align=PP_ALIGN.CENTER, font=HEADLINE_FONT, wrap=False)

# bottom pill bar (brand chip + subtitle)
pill_w = pw - 1.3
pill_x, pill_y, pill_h = px + 0.4, py + ph - 1.15, 0.85
add_rect(s, pill_x, pill_y, pill_w, pill_h, WHITE, radius=0.5)
add_text(s, pill_x + 0.3, pill_y, 2.6, pill_h, "202310318 안서연", 12.5, RED, bold=True,
          font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, pill_x + 3.05, pill_y + (pill_h - 0.42) / 2, 0.02, 0.42, MUTED)
add_text(s, pill_x + 3.3, pill_y, pill_w - 3.6, pill_h, "종류 · 특징 · 문화로 읽는 중국 소셜미디어 지형도",
          12, DARK, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# Slide 2 — Ecosystem overview
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, WHITE)
add_heading(s, "01 · OVERVIEW", "왜 중국은 SNS가 다를까?")
add_highlight_subtitle(s, "중국에는 만리방화벽(Great Firewall)이 만든 ", "독자적 소셜미디어", "가 있습니다.", 1.55)

rows = [
    ("Facebook", "facebook", "웨이보 (Weibo)", "weibo"),
    ("Instagram", "instagram", "샤오홍슈 (Xiaohongshu)", "xiaohongshu"),
    ("YouTube", "youtube", "빌리빌리 · 더우인", "bilibili"),
    ("Twitter (X)", "x_twitter", "웨이보 (Weibo)", "weibo"),
    ("KakaoTalk", "kakaotalk", "위챗 (WeChat)", "wechat"),
]
ry = 2.15
rh = 0.8
gap = 0.12
rx = 0.7
rw = 7.6
for i, (west, west_icon, china, icon_key) in enumerate(rows):
    y = ry + i * (rh + gap)
    add_rect(s, rx, y, rw, rh, GRAY_CARD, radius=0.12)
    add_icon_circle(s, west_icon, rx + 0.5, y + rh / 2, 0.5)
    add_text(s, rx + 0.85, y, 2.55, rh, west, 14, MUTED, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, rx + 3.45, y, 0.6, rh, "→", 20, GOLD, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    add_icon_circle(s, icon_key, rx + 4.55, y + rh / 2, 0.5)
    add_text(s, rx + 4.95, y, 2.5, rh, china, 14, DARK, bold=True, anchor=MSO_ANCHOR.MIDDLE)

# right stat card
scx, scy, scw, sch = 8.55, 2.15, 4.08, 4.48
add_rect(s, scx, scy, scw, sch, DARK_RED, radius=0.08)
add_text(s, scx, scy + 0.45, scw, 0.9, "10억+", 44, GLYPH_INK, bold=True, align=PP_ALIGN.CENTER, font=TITLE_FONT, wrap=False)
add_text(s, scx + 0.4, scy + 1.35, scw - 0.8, 0.7, "중국 모바일 인터넷 이용자 수", 13, LIGHT_TXT, align=PP_ALIGN.CENTER)
add_text(s, scx, scy + 2.45, scw, 0.9, "40+", 44, GLYPH_INK, bold=True, align=PP_ALIGN.CENTER, font=TITLE_FONT, wrap=False)
add_text(s, scx + 0.4, scy + 3.35, scw - 0.8, 0.9, "차단 이후 성장한\n자국산 주요 SNS 앱 수", 13, LIGHT_TXT, align=PP_ALIGN.CENTER)

# =====================================================================
# Slide 3 — Platforms: all 5 apps on one page
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, WHITE)
add_heading(s, "02 · PLATFORMS", "주요 플랫폼 — 5대 SNS")
add_highlight_subtitle(s, "종합 SNS부터 숏폼 · 라이프스타일까지, 중국을 대표하는 ", "5개 플랫폼", "을 소개합니다.", 1.55)

cards_all = [
    {
        "icon": "wechat", "name": "위챗", "cn": "WeChat · 微信",
        "tagline": "메신저 · 결제 · 미니프로그램의 슈퍼앱",
        "stat": "MAU 13억+",
    },
    {
        "icon": "weibo", "name": "웨이보", "cn": "Weibo · 微博",
        "tagline": "热搜 실시간 이슈로 여론을 움직이는 마이크로블로그",
        "stat": "热搜 랭킹",
    },
    {
        "icon": "douyin", "name": "더우인", "cn": "Douyin · 抖音",
        "tagline": "알고리즘 추천 기반 숏폼 · 라이브 커머스",
        "stat": "숏폼 · 커머스",
    },
    {
        "icon": "xiaohongshu", "name": "샤오홍슈", "cn": "Xiaohongshu · 小红书",
        "tagline": "뷰티 · 라이프스타일 후기 커뮤니티",
        "stat": "리뷰 · 커머스",
    },
    {
        "icon": "bilibili", "name": "빌리빌리", "cn": "Bilibili · 哔哩哔哩",
        "tagline": "탄막 문화가 살아있는 서브컬처 동영상",
        "stat": "탄막 · 팬덤",
    },
]
cy = 2.45
ch = 3.35
cw = 2.27
gap = 0.18
x0 = (SW - (5 * cw + 4 * gap)) / 2
cxs = [x0 + i * (cw + gap) for i in range(5)]
for cx, card in zip(cxs, cards_all):
    add_rect(s, cx, cy, cw, ch, GRAY_CARD, radius=0.09)
    add_icon_circle(s, card["icon"], cx + cw / 2, cy + 0.72, 0.95)
    add_text(s, cx + 0.1, cy + 1.28, cw - 0.2, 0.35, card["name"], 13.5, DARK, bold=True,
              align=PP_ALIGN.CENTER, font=TITLE_FONT, wrap=False)
    add_text(s, cx + 0.1, cy + 1.62, cw - 0.2, 0.3, card["cn"], 9, MUTED,
              align=PP_ALIGN.CENTER, wrap=False)
    add_text(s, cx + 0.17, cy + 2.0, cw - 0.34, 0.6, card["tagline"], 10.3, DARK, align=PP_ALIGN.CENTER)
    add_pill(s, cx + (cw - 1.95) / 2, cy + ch - 0.53, 1.95, 0.38, card["stat"], RED, WHITE, size=9.5)


# =====================================================================
# Slide 4 — Culture & conclusion
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, WHITE)
add_heading(s, "03 · CULTURE", "중국 SNS 문화의 특징")
add_highlight_subtitle(s, "서구 SNS와 단절된 중국은, 오히려 세계에서 가장 ",
                        "역동적이고 독자적인 SNS 생태계", "를 스스로 키워냈습니다.", 1.55)

cells = [
    ("网", "왕훙(网红) 경제", "인플루언서가 이끄는 소비 트렌드 · MCN 기획사의 산업화"),
    ("购", "라이브 커머스", "실시간 방송을 통한 직접 판매 · 왕훙 매출이 억대 규모로 성장"),
    ("审", "콘텐츠 검열", "민감 키워드 필터링 · 정부 규제에 따른 플랫폼 자체 심사"),
    ("梗", "밈(梗) 문화의 초고속 확산", "유행어 · 신조어가 폐쇄망 안에서 오히려 빠르고 응집력 있게 퍼짐"),
]
gx = 0.7
gy = 2.45
gw = 5.9
gh = 1.7
ggap = 0.25
positions = [
    (gx, gy), (gx + gw + ggap, gy),
    (gx, gy + gh + ggap), (gx + gw + ggap, gy + gh + ggap),
]
for (cx0, cy0), (glyph, title, desc) in zip(positions, cells):
    add_rect(s, cx0, cy0, gw, gh, GRAY_CARD, radius=0.08)
    add_badge(s, cx0 + 0.75, cy0 + gh / 2, 0.85, glyph, 26)
    add_text(s, cx0 + 1.35, cy0 + 0.25, gw - 1.6, 0.45, title, 15.5, RED, bold=True, font=TITLE_FONT)
    add_text(s, cx0 + 1.35, cy0 + 0.72, gw - 1.6, gh - 0.9, desc, 12, DARK)


# =====================================================================
# Slide 5 — Thank you
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, RED)

add_text(s, 0.85, 0.9, 11.6, 0.4, "THANK YOU", 14, GLYPH_INK, bold=True, align=PP_ALIGN.CENTER, font=BODY_FONT, wrap=False)
add_text(s, 0.85, SH / 2 - 1.35, 11.6, 1.6, "감사합니다", 60, WHITE, bold=False,
          align=PP_ALIGN.CENTER, font=HEADLINE_FONT, wrap=False)
add_text(s, 0.85, SH / 2 + 0.35, 11.6, 0.5, "중국사회문화 캡스톤디자인 · 중국 SNS 완전정복", 14, LIGHT_TXT,
          align=PP_ALIGN.CENTER, font=BODY_FONT, wrap=False)

out_path = "china_sns_deck.pptx"
prs.save(out_path)
print("saved", out_path)
