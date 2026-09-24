# -*- coding: utf-8 -*-
import os
import re
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# PUBLIC=1 builds a contact-info-redacted copy for publishing (e.g. GitHub).
PUBLIC = os.environ.get("PUBLIC") == "1"
BIRTHDATE  = "20XX. XX. XX"      if PUBLIC else "2004. 05. 13"
ADDRESS    = "경기도 OO시"        if PUBLIC else "경기도 동두천시"
PHONE      = "010-OOOO-OOOO"     if PUBLIC else "010-7940-0513"
EMAIL      = "OOO@email.com"     if PUBLIC else "aseo13yeon@gmail.com"

ASSET_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "portfolio")
SKILL_DIR = os.path.join(ASSET_DIR, "skills")
TOOL_DIR = os.path.join(ASSET_DIR, "tools")
IMG = {
    "profile": os.path.join(ASSET_DIR, "profile.png"),
    "p1_vocab": os.path.join(ASSET_DIR, "project1_vocab.png"),
    "p1_flashcard": os.path.join(ASSET_DIR, "project1_flashcard.png"),
    "p2_cover": os.path.join(ASSET_DIR, "project2_cover.png"),
    "p2_platforms": os.path.join(ASSET_DIR, "project2_platforms.png"),
    "skill_ps": os.path.join(SKILL_DIR, "ps.png"),
    "skill_ai": os.path.join(SKILL_DIR, "ai.png"),
    "skill_figma": os.path.join(SKILL_DIR, "figma.png"),
    "skill_excel": os.path.join(SKILL_DIR, "excel.png"),
    "skill_ppt": os.path.join(SKILL_DIR, "ppt.png"),
    "skill_notion": os.path.join(SKILL_DIR, "notion.png"),
    "skill_chatgpt": os.path.join(SKILL_DIR, "chatgpt.png"),
    "skill_gemini": os.path.join(SKILL_DIR, "gemini.png"),
    "tool_html": os.path.join(TOOL_DIR, "html.png"),
    "tool_css": os.path.join(TOOL_DIR, "css.png"),
    "tool_js": os.path.join(TOOL_DIR, "js.png"),
    "tool_ppt": os.path.join(TOOL_DIR, "ppt.png"),
    "tool_python": os.path.join(TOOL_DIR, "python.png"),
}

# ---------- palette (only these text colors are used) ----------
BG          = RGBColor(0xF7, 0xF7, 0xF5)  # portfolio background
ACCENT      = RGBColor(0x7F, 0x9B, 0xC7)  # title text color
INK         = RGBColor(0x37, 0x37, 0x3B)  # body text color
ACCENT_TINT = RGBColor(0xF0, 0xF3, 0xF8)  # structural fill only (not text)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
LINE        = RGBColor(0xE1, 0xE1, 0xDE)  # structural border only

TITLE_FONT = "A2Z 7 Bold"     # 제목
BODY_FONT  = "A2Z 4 Regular"  # 본문
CJK_FONT   = "Malgun Gothic"  # fallback for Chinese hanzi glyphs not in A2Z

SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width = Emu(int(SW * 914400))
prs.slide_height = Emu(int(SH * 914400))
BLANK = prs.slide_layouts[6]

HANZI_RE = re.compile(r'([一-鿿]+)')


def no_shadow(shp):
    shp.shadow.inherit = False


def add_rect(slide, x, y, w, h, fill, radius=None, line_color=None, line_w=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_color is not None:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(line_w or 0.75)
    else:
        shp.line.fill.background()
    no_shadow(shp)
    return shp


def add_run_mixed(p, text, size, color, font, bold=False):
    for seg in [s for s in HANZI_RE.split(text) if s]:
        r = p.add_run()
        r.text = seg
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = CJK_FONT if HANZI_RE.fullmatch(seg) else font
        r.font.color.rgb = color


def add_text(slide, x, y, w, h, text, size, color, font=BODY_FONT, bold=False,
             align=PP_ALIGN.LEFT, anchor=None, wrap=True, line_spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if anchor is not None:
        tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = line_spacing
        add_run_mixed(p, line, size, color, font, bold=bold)
    return tb


def add_heading(slide, x, y, w, h, text, size, align=PP_ALIGN.LEFT, anchor=None, wrap=True, line_spacing=1.0):
    return add_text(slide, x, y, w, h, text, size, ACCENT, font=TITLE_FONT, bold=True,
                     align=align, anchor=anchor, wrap=wrap, line_spacing=line_spacing)


def add_body(slide, x, y, w, h, text, size, align=PP_ALIGN.LEFT, anchor=None, wrap=True, line_spacing=1.15):
    return add_text(slide, x, y, w, h, text, size, INK, font=BODY_FONT, bold=False,
                     align=align, anchor=anchor, wrap=wrap, line_spacing=line_spacing)


def add_bullets(slide, x, y, w, h, items, size, space_after=8, line_spacing=1.15):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        r1 = p.add_run()
        r1.text = "•  "
        r1.font.size = Pt(size)
        r1.font.name = BODY_FONT
        r1.font.color.rgb = ACCENT
        add_run_mixed(p, item, size, INK, BODY_FONT)
    return tb


def add_badge(slide, x, y, d, label, size=9):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.adjustments[0] = 0.28
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = LINE
    shp.line.width = Pt(1)
    no_shadow(shp)
    add_text(slide, x, y, d, d, label, size, INK, font=BODY_FONT,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    return shp


def add_icon_badge(slide, x, y, d, image_key, pad=0.07):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.adjustments[0] = 0.28
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = LINE
    shp.line.width = Pt(1)
    no_shadow(shp)
    picture_fit(slide, IMG[image_key], x + pad, y + pad, d - 2 * pad, d - 2 * pad)
    return shp


def picture_cover(slide, path, x, y, box_w, box_h):
    iw, ih = Image.open(path).size
    img_ratio = iw / ih
    box_ratio = box_w / box_h
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(box_w), Inches(box_h))
    if img_ratio > box_ratio:
        visible = box_ratio / img_ratio
        crop = (1 - visible) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        visible = img_ratio / box_ratio
        crop = (1 - visible) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    return pic


def picture_fit(slide, path, x, y, box_w, box_h):
    iw, ih = Image.open(path).size
    ratio = iw / ih
    box_ratio = box_w / box_h
    if ratio > box_ratio:
        w = box_w
        h = w / ratio
    else:
        h = box_h
        w = h * ratio
    px = x + (box_w - w) / 2
    py = y + (box_h - h) / 2
    return slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(w), Inches(h))


def add_page_footer(slide, n):
    add_text(slide, SW - 1.1, SH - 0.5, 0.75, 0.3, f"{n} / 5", 10, INK,
              align=PP_ALIGN.RIGHT, wrap=False)


def add_project_header(slide, kicker, category, title, period, tools, title_size=27):
    add_rect(slide, 0.65, 0.44, 2.05, 0.36, ACCENT, radius=0.5)
    add_text(slide, 0.65, 0.44, 2.05, 0.36, kicker, 12, WHITE, wrap=False,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_body(slide, 0.75, 0.98, 6.0, 0.35, category, 12.5)
    add_heading(slide, 0.72, 1.34, 6.3, 0.85, title, title_size, wrap=False, anchor=MSO_ANCHOR.TOP)
    add_body(slide, 0.75, 1.91, 6.0, 0.35, period, 12.5)
    bx = 0.75
    for t in tools:
        add_icon_badge(slide, bx, 2.34, 0.5, t)
        bx += 0.62


# =====================================================================
# Slide 1 — Cover
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, BG)

add_body(s, 0.75, 0.55, 4.0, 0.3, "2026 PORTFOLIO", 11)
add_text(s, SW - 4.75, 0.55, 4.0, 0.3, "AN SEO YEON", 11, INK, font=BODY_FONT, align=PP_ALIGN.RIGHT, wrap=False)

add_heading(s, 0, 2.75, SW, 1.5, "PORTFOLIO", 90, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)

add_rect(s, SW / 2 - 0.8, 4.62, 1.6, 0.028, ACCENT)

add_body(s, 0, 4.95, SW, 0.5, "안서연 · An Seo Yeon", 20, align=PP_ALIGN.CENTER, wrap=False)
add_body(s, 1.5, 5.5, SW - 3.0, 0.5, "브랜드와 소비자 사이, 콘텐츠로 결을 만듭니다.", 13.5, align=PP_ALIGN.CENTER)

# =====================================================================
# Slide 2 — Self introduction
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, BG)
add_rect(s, 0, 0, 4.6, SH, ACCENT_TINT)

add_heading(s, 0.55, 0.55, 3.7, 1.3, "안녕하세요.\n단어 하나로 브랜드의 결을\n완성하는 안서연입니다.", 18.5, line_spacing=1.15)

# profile photo with a thin frame (3:4 ID-photo ratio)
photo_w, photo_h = 2.0625, 2.75
photo_x, photo_y = 0.55, 2.35
add_rect(s, photo_x - 0.03, photo_y - 0.03, photo_w + 0.06, photo_h + 0.06, WHITE, line_color=LINE, line_w=1)
s.shapes.add_picture(IMG["profile"], Inches(photo_x), Inches(photo_y), Inches(photo_w), Inches(photo_h))

add_heading(s, 0.55, 5.35, 3.7, 0.4, "안서연 / An Seo Yeon", 17)
add_body(s, 0.55, 5.78, 3.7, 0.3, BIRTHDATE, 11)
add_body(s, 0.55, 6.18, 3.7, 0.9, f"{ADDRESS}\nTel. {PHONE}\nEmail. {EMAIL}", 11, line_spacing=1.35)

# ---- Column A: EDUCATION / EXPERIENCE ----
ax = 5.0
add_heading(s, ax, 0.55, 3.6, 0.35, "EDUCATION", 15)
add_body(s, ax, 1.0, 0.85, 0.5, "2023. 03", 10.5)
add_body(s, ax + 0.95, 1.0, 2.7, 0.6, "OO대학교 중어중문전공\nOO대학교 경영학전공 복수전공", 11, line_spacing=1.3)

add_heading(s, ax, 2.05, 3.6, 0.35, "EXPERIENCE", 15)

experience = [
    ("중어중문전공 제36대 OO학생회 홍보부원", "2024. 03 ~ 2024. 12"),
    ("OO기관 제11차 한일중 대학생 외교캠프 한국팀원", "2024. 07 ~ 2024. 07"),
    ("OO기업 헤이영 서포터즈 3기 프로젝트 팀장", "2025. 08 ~ 2025. 12"),
    ("OO대학교 마케팅학회 OO", "2026. 03 ~ 2027. 01"),
    ("OO 펀딩 프로젝트", "2025. 12 ~ 2026. 02"),
    ("OO기업 그린핑거 AI CREW 1기", "2026. 08 ~ 2026. 11"),
    ("전통시장 온라인 마케팅 프로젝트", "2025. 06 ~ 2026. 07"),
]
ey = 2.48
for title, period in experience:
    add_body(s, ax, ey, 3.6, 0.3, title, 10.8)
    add_body(s, ax, ey + 0.27, 3.6, 0.25, period, 9)
    ey += 0.615

# ---- Column B: AWARDS / CERTIFICATION / SKILL ----
bx = 9.1
add_heading(s, bx, 0.55, 3.7, 0.35, "AWARDS", 15)
awards = [
    ("[OO기관] 제11차 한일중 대학생 외교캠프 우수발표상", "2024. 07"),
    ("[OO기업] 헤이영 서포터즈 3기 우수 서포터즈", "2025. 12"),
]
ay = 1.0
for title, period in awards:
    add_body(s, bx, ay, 3.7, 0.3, title, 10.8)
    add_body(s, bx, ay + 0.27, 3.7, 0.25, period, 9)
    ay += 0.615

add_heading(s, bx, 2.55, 3.7, 0.35, "CERTIFICATION", 15)
certs = [
    ("2024. 08", "컴퓨터활용능력 2급"),
    ("2025. 10", "GTQ 1급"),
    ("2025. 11", "GTQi 1급"),
    ("2026. 05", "HSK 6급"),
    ("2026. 05", "데이터분석준전문가(ADsP)"),
]
cy = 3.0
for period, name in certs:
    add_body(s, bx, cy, 0.85, 0.3, period, 10.5)
    add_body(s, bx + 0.95, cy, 2.75, 0.3, name, 11)
    cy += 0.36

add_heading(s, bx, 5.15, 3.7, 0.35, "SKILL", 15)
skills = ["skill_ps", "skill_ai", "skill_figma", "skill_excel",
          "skill_ppt", "skill_notion", "skill_chatgpt", "skill_gemini"]
sx = bx
sy = 5.6
for i, sk in enumerate(skills):
    if i == 4:
        sx = bx
        sy += 0.62
    add_icon_badge(s, sx, sy, 0.5, sk)
    sx += 0.62

add_page_footer(s, 2)

# =====================================================================
# Slide 3 — Main Project 01 (capstone website)
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, BG)

add_project_header(
    s,
    kicker="MAIN PROJECT 01",
    category="개인 프로젝트 · 웹사이트 제작",
    title="중국어 학습 웹사이트 — 生词本 & 汉字卡片",
    period="2026. 09  ·  중국사회문화 캡스톤디자인",
    tools=["tool_html", "tool_css", "tool_js"],
    title_size=23,
)

add_heading(s, 0.72, 2.95, 5.8, 0.35, "배경 및 목적", 15)
add_body(s, 0.75, 3.35, 5.9, 1.05,
         "HSK 어휘와 한자를 반복해서 암기할 수 있는 도구가 마땅치 않아, "
         "중국어 전공자로서 직접 사용할 학습 웹사이트를 기획·제작함. "
         "실제 학습 과정에 필요한 기능을 스스로 설계해보는 것을 목표로 함.",
         11.5, line_spacing=1.35)

add_heading(s, 0.72, 4.5, 5.8, 0.35, "실행방안", 15)
add_bullets(s, 0.75, 4.9, 5.9, 1.5, [
    "生词本(단어장): 한자·병음·뜻 등록과 검색, 암기 체크, 유의어 2지선다 퀴즈, "
    "발음이 비슷하거나 성조만 다른 단어를 자동 감지해 보여주는 기능을 구현",
    "汉字卡片(플래시카드): HSK 1~2급 단어를 카드로 넘기며 아는/모르는 단어를 표시하고, "
    "즐겨찾기·순서 섞기로 반복 학습이 가능하도록 구현",
], 11, space_after=8, line_spacing=1.3)

add_heading(s, 0.72, 6.15, 5.8, 0.35, "결과", 15)
add_body(s, 0.75, 6.52, 5.9, 0.3, "HTML/CSS/JavaScript로 학습 도구 2종을 직접 구현, 개인 HSK 학습에 실제 활용 중", 11.5)

add_body(s, 7.0, 1.0, 5.6, 0.3, "▼ 生词本 · 汉字卡片 (웹 브라우저 화면)", 11.5)
frame_w, frame_h = 2.75, 5.38
fx1, fx2, fy = 7.0, 10.0, 1.4
add_rect(s, fx1, fy, frame_w, frame_h, WHITE, radius=0.06, line_color=LINE, line_w=1.25)
picture_cover(s, IMG["p1_vocab"], fx1 + 0.15, fy + 0.15, frame_w - 0.3, frame_h - 0.3)
add_rect(s, fx2, fy, frame_w, frame_h, WHITE, radius=0.06, line_color=LINE, line_w=1.25)
picture_cover(s, IMG["p1_flashcard"], fx2 + 0.15, fy + 0.15, frame_w - 0.3, frame_h - 0.3)

add_page_footer(s, 3)

# =====================================================================
# Slide 4 — Main Project 02 (SNS deck)
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, BG)

add_project_header(
    s,
    kicker="MAIN PROJECT 02",
    category="팀 발표 프로젝트 · 프레젠테이션 제작",
    title="중국 SNS 완전정복",
    period="2026. 09  ·  중국사회문화 캡스톤디자인 4주차",
    tools=["tool_ppt", "tool_python"],
)

add_heading(s, 0.72, 2.95, 5.8, 0.35, "배경 및 목적", 15)
add_body(s, 0.75, 3.35, 5.9, 1.05,
         "만리방화벽(Great Firewall) 이후 독자적으로 발전한 중국 SNS 생태계를, "
         "서구 SNS와의 비교와 문화적 맥락을 통해 쉽게 전달하기 위해 발표자료를 제작함.",
         11.5, line_spacing=1.35)

add_heading(s, 0.72, 4.15, 5.8, 0.35, "실행방안", 15)
add_bullets(s, 0.75, 4.55, 5.9, 1.85, [
    "서구 SNS와 중국 대체 플랫폼을 나란히 비교해 생태계 구조를 한눈에 전달",
    "위챗·웨이보·더우인·샤오홍슈·빌리빌리 5대 플랫폼을 카드형으로 정리",
    "왕훙 경제·라이브커머스·콘텐츠 검열·밈 문화, 4가지 키워드로 마무리",
    "python-pptx로 슬라이드를 코드 기반 제작해 디자인 재현성과 수정 효율을 확보",
], 10.8, space_after=6, line_spacing=1.25)

add_heading(s, 0.72, 6.15, 5.8, 0.35, "결과", 15)
add_body(s, 0.75, 6.52, 5.9, 0.3, "5페이지 발표자료 완성, 중국사회문화 캡스톤디자인 수업 발표 진행", 11.5)

add_body(s, 7.4, 1.0, 5.6, 0.3, "▼ 중국 SNS 완전정복 (발표자료 일부)", 11.5)
iw, ih = 4.75, 2.672
ix, iy1 = 7.4, 1.42
add_rect(s, ix, iy1, iw, ih, WHITE, radius=0.04, line_color=LINE, line_w=1.25)
picture_fit(s, IMG["p2_cover"], ix + 0.1, iy1 + 0.1, iw - 0.2, ih - 0.2)
iy2 = iy1 + ih + 0.3
add_rect(s, ix, iy2, iw, ih, WHITE, radius=0.04, line_color=LINE, line_w=1.25)
picture_fit(s, IMG["p2_platforms"], ix + 0.1, iy2 + 0.1, iw - 0.2, ih - 0.2)

add_page_footer(s, 4)

# =====================================================================
# Slide 5 — Closing
# =====================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, BG)

add_heading(s, 0, 2.85, SW, 1.1, "감사합니다", 56, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
add_rect(s, SW / 2 - 0.8, 4.15, 1.6, 0.028, ACCENT)
add_body(s, 1.5, 4.45, SW - 3.0, 0.6, "브랜드와 소비자 사이, 다음 이야기를 계속 만들어가겠습니다.", 14.5, align=PP_ALIGN.CENTER)
add_body(s, 0, 5.05, SW, 0.3, f"안서연 · {EMAIL} · {PHONE}", 11.5, align=PP_ALIGN.CENTER, wrap=False)

add_page_footer(s, 5)

out_name = "portfolio_deck_public.pptx" if PUBLIC else "portfolio_deck.pptx"
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), out_name)
prs.save(out_path)
print("saved", out_path)
